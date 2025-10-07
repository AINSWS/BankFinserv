#!/usr/bin/env python3
"""Direct test of the formatting function"""

import pandas as pd
import sys
import os
sys.path.insert(0, 'src')

# Create test data
test_data = pd.DataFrame({
    'loan_id': [181669, 181670, 181671, 181672],
    'customer_name': ['VAISHALI SANTOSH GAYAKWAD', 'Customer 2', 'Customer 3', 'Customer 4'],
    'status': ['MATCHED - Perfect Match', 'MATCHED - Phase 3', 'MATCHED - Group Payment', 'MISMATCH - Amount'],
    'system_amount': [6565.0, 3200.0, 2730.0, 1800.0],
    'qr_amount': [6565.0, 3200.0, 2730.0, 0.0],
    'branch': ['CHIPLUN', 'MUMBAI', 'PUNE', 'DELHI'],
    'group_name': ['CHI-33-CHI SUNITA SATISH GAYAKWAD', 'Group 2', 'Group 3', 'Group 4']
})

summary = {
    'total_unique_loans': 4,
    'perfect_matches': 1,
    'minor_matches': 0,
    'total_matches': 3,
    'amount_mismatches': 1,
    'match_percentage': 75.0
}

# Export manually using same logic as UI
file_path = "direct_test_export.xlsx"

print("📝 Creating Excel file with manual export logic...")

with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
    # Find status column
    status_column = None
    for col in ['status', 'reconciliation_status']:
        if col in test_data.columns:
            status_column = col
            break
    
    if status_column and not test_data.empty:
        # Separate by status
        matched_df = test_data[test_data[status_column].str.contains('MATCHED', na=False)]
        mismatch_df = test_data[test_data[status_column].str.contains('MISMATCH', na=False)]
        
        # Export matched records
        if not matched_df.empty:
            matched_df.to_excel(writer, sheet_name='Matched_Records', index=False)
            print(f"   ✓ Matched records: {len(matched_df)} entries")
        
        # Export mismatched records  
        if not mismatch_df.empty:
            mismatch_df.to_excel(writer, sheet_name='Mismatched_Records', index=False)
            print(f"   ✓ Mismatched records: {len(mismatch_df)} entries")
        
        # Export all records for reference
        test_data.to_excel(writer, sheet_name='All_Records', index=False)
        print(f"   ✓ All records: {len(test_data)} entries")
    
    # Create summary sheet
    summary_data = {
        'Metric': ['Total Records', 'Perfect Matches', 'Minor Matches', 'Total Matches', 'Amount Mismatches', 'Match Rate %'],
        'Value': [
            summary.get('total_unique_loans', 0),
            summary.get('perfect_matches', 0),
            summary.get('minor_matches', 0),
            summary.get('total_matches', 0),
            summary.get('amount_mismatches', 0),
            f"{summary.get('match_percentage', 0):.1f}%"
        ]
    }
    summary_df = pd.DataFrame(summary_data)
    summary_df.to_excel(writer, sheet_name='Summary', index=False)
    print(f"   ✓ Summary created")

# Now apply the formatting function
print(f"\n🎨 Applying formatting...")

try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font
    
    # Load the workbook
    print(f"📂 Loading workbook...")
    wb = load_workbook(file_path)
    print(f"📋 Sheets found: {wb.sheetnames}")
    
    # Define clean colors
    matched_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Light green
    mismatch_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Light red
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")   # Light gray
    header_font = Font(bold=True)
    
    for sheet_name in wb.sheetnames:
        print(f"🎨 Processing sheet: {sheet_name}")
        ws = wb[sheet_name]
        
        # Apply auto-filter to make it visible
        if ws.max_row > 1 and ws.max_column > 1:
            filter_ref = f"A1:{ws.cell(row=1, column=ws.max_column).coordinate}{ws.max_row}"
            ws.auto_filter.ref = filter_ref
            print(f"   ✅ Auto-filter applied: {filter_ref}")
        
        # Format headers
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
        print(f"   🎨 Headers formatted (bold + gray background)")
        
        # Apply row coloring based on sheet name and status
        if 'Matched' in sheet_name:
            # Light green for matched records
            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = matched_fill
            print(f"   🟢 Applied green background to {ws.max_row-1} rows")
                    
        elif 'Mismatched' in sheet_name:
            # Light red for mismatched records  
            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = mismatch_fill
            print(f"   🔴 Applied red background to {ws.max_row-1} rows")
        
        elif 'All' in sheet_name:
            # Apply conditional coloring based on status column
            status_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value and 'status' in str(ws.cell(row=1, column=col).value).lower():
                    status_col = col
                    break
            
            if status_col:
                print(f"   🔍 Status column found at column {status_col}")
                for row in range(2, ws.max_row + 1):
                    status_value = ws.cell(row=row, column=status_col).value
                    if status_value and 'MATCHED' in str(status_value):
                        for col in range(1, ws.max_column + 1):
                            ws.cell(row=row, column=col).fill = matched_fill
                        print(f"   🟢 Row {row} colored green (matched)")
                    elif status_value and 'MISMATCH' in str(status_value):
                        for col in range(1, ws.max_column + 1):
                            ws.cell(row=row, column=col).fill = mismatch_fill
                        print(f"   🔴 Row {row} colored red (mismatch)")
        
        # Auto-adjust column widths for better visibility
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)  # Cap at 30 chars
            ws.column_dimensions[column_letter].width = adjusted_width
        print(f"   📏 Column widths auto-adjusted")
    
    # Save the formatted workbook
    wb.save(file_path)
    print(f"✅ Clean formatting applied successfully!")
    print(f"📁 Check the file: {file_path}")
    
    # Verify sheets
    excel_file = pd.ExcelFile(file_path)
    print(f"\n📋 Final verification - Sheets: {excel_file.sheet_names}")
    for sheet in excel_file.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet)
        print(f"   • {sheet}: {len(df)} rows")
    
    print(f"\n✨ Features applied:")
    print(f"   ✅ Auto-filters on all sheets (dropdown arrows in headers)")
    print(f"   ✅ Color coding: Green=Matched, Red=Mismatched")
    print(f"   ✅ Bold gray headers")
    print(f"   ✅ Auto-sized columns")
    
except Exception as e:
    print(f"❌ Formatting failed: {e}")
    import traceback
    traceback.print_exc()