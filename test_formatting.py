#!/usr/bin/env python3
"""Test formatting function to ensure it works"""

import pandas as pd
import sys
import os
sys.path.insert(0, 'src')

# Create test data
test_data = pd.DataFrame({
    'loan_id': [181669, 181670, 181671],
    'customer_name': ['Customer 1', 'Customer 2', 'Customer 3'],
    'status': ['MATCHED - Perfect Match', 'MATCHED - Phase 3', 'MISMATCH - Amount'],
    'amount': [1000.0, 2000.0, 3000.0]
})

# Export to Excel
test_file = 'test_format.xlsx'
with pd.ExcelWriter(test_file, engine='openpyxl') as writer:
    test_data.to_excel(writer, sheet_name='Test_Data', index=False)

print(f"📝 Test file created: {test_file}")

# Now test the formatting function
try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font
    
    print("🔧 Applying formatting...")
    
    # Load the workbook
    wb = load_workbook(test_file)
    ws = wb['Test_Data']
    
    # Define colors
    matched_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Light green
    mismatch_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Light red
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")   # Light gray
    header_font = Font(bold=True)
    
    print(f"   📊 Sheet has {ws.max_row} rows and {ws.max_column} columns")
    
    # Apply auto-filter
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=ws.max_column).coordinate}{ws.max_row}"
        print(f"   ✅ Auto-filter applied: A1:{ws.cell(row=1, column=ws.max_column).coordinate}{ws.max_row}")
    
    # Format headers
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        print(f"   🎨 Header formatted: {cell.coordinate} = {cell.value}")
    
    # Find status column
    status_col = None
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value and 'status' in str(ws.cell(row=1, column=col).value).lower():
            status_col = col
            print(f"   🔍 Status column found at column {col}")
            break
    
    # Apply row coloring
    if status_col:
        for row in range(2, ws.max_row + 1):
            status_value = ws.cell(row=row, column=status_col).value
            print(f"   📝 Row {row} status: {status_value}")
            
            if status_value and 'MATCHED' in str(status_value):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = matched_fill
                print(f"   🟢 Row {row} colored green (matched)")
            elif status_value and 'MISMATCH' in str(status_value):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = mismatch_fill
                print(f"   🔴 Row {row} colored red (mismatch)")
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column_letter].width = adjusted_width
        print(f"   📏 Column {column_letter} width set to {adjusted_width}")
    
    # Save
    wb.save(test_file)
    print(f"✅ Formatting applied successfully!")
    print(f"📁 Check the file: {test_file}")
    
except Exception as e:
    print(f"❌ Formatting failed: {e}")
    import traceback
    traceback.print_exc()