#!/usr/bin/env python3
"""
Quick test to verify that Excel files have visible formatting
Open the generated Excel file and check for:
1. Dropdown arrows in headers (auto-filters)
2. Green background on matched records
3. Red background on mismatched records
4. Bold gray headers
"""

import pandas as pd
import os

# Create sample data
data = pd.DataFrame({
    'loan_id': [181669, 181670, 181671],
    'customer_name': ['Customer A', 'Customer B', 'Customer C'],
    'status': ['MATCHED - Perfect', 'MATCHED - Phase 3', 'MISMATCH - Amount'],
    'amount': [1000, 2000, 3000]
})

# Export to Excel
test_file = 'formatting_test.xlsx'

# Basic export
with pd.ExcelWriter(test_file, engine='openpyxl') as writer:
    # Matched records
    matched = data[data['status'].str.contains('MATCHED', na=False)]
    matched.to_excel(writer, sheet_name='Matched_Records', index=False)
    
    # Mismatched records
    mismatched = data[data['status'].str.contains('MISMATCH', na=False)]
    mismatched.to_excel(writer, sheet_name='Mismatched_Records', index=False)

print(f"📝 Basic file created: {test_file}")

# Apply formatting
try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font
    
    wb = load_workbook(test_file)
    
    # Colors
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    bold_font = Font(bold=True)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Auto-filter (creates dropdown arrows)
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{ws.cell(1, ws.max_column).coordinate}{ws.max_row}"
            print(f"✅ Auto-filter added to {sheet_name}")
        
        # Format headers (bold + gray)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(1, col)
            cell.font = bold_font
            cell.fill = header_fill
        
        # Color rows based on content
        if 'Matched' in sheet_name:
            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row, col).fill = green_fill
            print(f"🟢 Green background applied to {sheet_name}")
            
        elif 'Mismatched' in sheet_name:
            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row, col).fill = red_fill
            print(f"🔴 Red background applied to {sheet_name}")
        
        # Auto-size columns
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 25)
    
    wb.save(test_file)
    print(f"✅ Formatting applied successfully!")
    
    print(f"\n📁 Open this file to verify formatting: {test_file}")
    print(f"🔍 Look for:")
    print(f"   • Dropdown arrows in header row (auto-filters)")
    print(f"   • Green background on Matched_Records sheet")
    print(f"   • Red background on Mismatched_Records sheet")
    print(f"   • Bold gray headers")
    print(f"   • Properly sized columns")
    
except Exception as e:
    print(f"❌ Formatting error: {e}")
    import traceback
    traceback.print_exc()